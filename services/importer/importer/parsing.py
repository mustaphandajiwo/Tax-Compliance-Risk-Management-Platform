"""CSV / XLSX parsing into string rows. No type coercion happens here."""

from __future__ import annotations

import csv
import io

from openpyxl import load_workbook

MAX_ROWS = 500_000


class ParseError(Exception):
    pass


def _clean(v: object) -> str:
    if v is None:
        return ""
    if hasattr(v, "isoformat"):
        return v.isoformat()  # type: ignore[no-any-return]
    return str(v).strip()


def parse_file(filename: str, content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    name = filename.lower()
    if name.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ParseError("CSV must be UTF-8 encoded") from None
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise ParseError("CSV has no header row")
        headers = [h.strip() for h in reader.fieldnames]
        if len(set(headers)) != len(headers):
            raise ParseError("CSV header has duplicate column names")
        rows = []
        for i, row in enumerate(reader):
            if i >= MAX_ROWS:
                raise ParseError(f"file exceeds {MAX_ROWS} rows")
            if None in row:
                raise ParseError(f"row {i + 2} has more values than header columns")
            rows.append({k.strip(): _clean(v) for k, v in row.items()})
        return headers, rows
    if name.endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise ParseError(f"cannot read XLSX: {exc}") from None
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        try:
            header_row = next(it)
        except StopIteration:
            raise ParseError("XLSX sheet is empty") from None
        headers = [_clean(h) for h in header_row]
        if len(set(headers)) != len(headers):
            raise ParseError("XLSX header has duplicate column names")
        rows = []
        for i, values in enumerate(it):
            if i >= MAX_ROWS:
                raise ParseError(f"file exceeds {MAX_ROWS} rows")
            if all(v is None for v in values):
                continue
            rows.append({h: _clean(v) for h, v in zip(headers, values, strict=False)})
        return headers, rows
    raise ParseError("only .csv and .xlsx files are accepted")
